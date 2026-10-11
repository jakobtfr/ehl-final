"""Reconcile local BMW inputs against actual bmw-import-v1 results without emitting source text.

Run against a fresh, ignored database for first-import counts. A second import proves that the
same original bytes do not multiply evidence. Output contains file names, hashes, locations and
counts only; no feedback, document text, or spreadsheet cell values.
"""

import argparse
import hashlib
import io
import json
import os
import sys
from collections import Counter
from contextlib import closing
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from database.store import connect, init_database
from server.importers import file_scopes, parse_scopes
from server.market_data import SUPPORTED_SUFFIXES
from server.platform import run_capability

METHOD = "bmw-source-reconciliation-v1"


def _surfaces(path: Path, data: bytes, located: dict) -> dict:
    """List sheets/pages with no parsed evidence; structural rows are not evidence."""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            names = workbook.sheetnames
        finally:
            workbook.close()
        used = {item["sheet"] for item in located.values() if "sheet" in item}
        return {"sheets": len(names), "ignored_sheets": [name for name in names if name not in used]}
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader

        pages = len(PdfReader(io.BytesIO(data)).pages)
        used = {item["page"] for item in located.values() if "page" in item}
        return {"pages": pages, "ignored_pages": [page for page in range(1, pages + 1) if page not in used]}
    return {}


def reconcile(folder: Path, first: dict, repeat: dict, manifest: dict, connection) -> dict:
    """Compare the API's per-scope result with the parser and database, using only redacted fields."""
    expected = {
        Path(item["path"]).name: item
        for item in manifest["files"]
        if item["kind"] in {"google-drive-input", "external-document"}
    }
    first_files = {Path(item["path"]).name: item for item in first["files"]}
    repeat_files = {Path(item["path"]).name: item for item in repeat["files"]}
    files = []
    for path in sorted(folder.iterdir()):
        if not path.is_file() or path.name.startswith(".") or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        data = path.read_bytes()
        digest = hashlib.sha256(data).hexdigest()
        outcome = first_files[path.name]
        again = repeat_files[path.name]
        scopes = [(item["vehicle"], item["market"]) for item in outcome["scopes"]]
        parsed = parse_scopes(path.name, data, scopes)
        declared = file_scopes(path.name, data)
        # A locator may be repeated across scopes; count a sheet/page as used once.
        located = {
            json.dumps(record["location"], sort_keys=True): record["location"]
            for records in parsed.values()
            for record in records
        }
        scope_rows = []
        for item in outcome["scopes"]:
            key = (item["vehicle"], item["market"])
            source = connection.execute(
                "SELECT id, content_hash, original_path FROM sources WHERE id=?", (item["source_id"],)
            ).fetchone()
            if source is None or source["content_hash"] != digest:
                raise AssertionError(f"Persisted source hash mismatch: {path.name} {key}")
            rows = connection.execute(
                "SELECT location, metadata, observed_at FROM evidence WHERE source_id=?", (item["source_id"],)
            ).fetchall()
            stored_locations = Counter(json.dumps(json.loads(row["location"]), sort_keys=True) for row in rows)
            parsed_locations = Counter(json.dumps(record["location"], sort_keys=True) for record in parsed[key])
            if stored_locations - parsed_locations:
                raise AssertionError(f"Persisted locator is absent from parser output: {path.name} {key}")
            statuses = Counter(json.loads(row["metadata"]).get("date_status", "unknown") for row in rows)
            record_types = Counter(json.loads(row["metadata"]).get("record_type", "unknown") for row in rows)
            if item["record_count"] != len(rows):
                raise AssertionError(f"Import count differs from persisted evidence: {path.name} {key}")
            repeat_scope = next((row for row in again["scopes"] if (row["vehicle"], row["market"]) == key), None)
            if (
                repeat_scope is None
                or not repeat_scope["deduplicated"]
                or repeat_scope["source_id"] != item["source_id"]
            ):
                raise AssertionError(f"Repeat import did not deduplicate: {path.name} {key}")
            scope_rows.append(
                {
                    "vehicle": key[0],
                    "market": key[1],
                    "assignment": (
                        "source"
                        if any(d["vehicle"] == key[0] and d["market"] == key[1] for d in declared)
                        else "expanded_from_source"
                        if any(d["vehicle"] == key[0] and d["market"] is None for d in declared)
                        else "document_default"
                    ),
                    "source_id": item["source_id"],
                    "original_path": source["original_path"],
                    "parsed_records": len(parsed[key]),
                    "stored_records": len(rows),
                    "same_source_duplicate_records": len(parsed[key]) - len(rows),
                    "deduplicated_on_first_pass": item["deduplicated"],
                    "deduplicated_on_repeat": repeat_scope["deduplicated"],
                    "record_types": dict(sorted(record_types.items())),
                    "date_status": dict(sorted(statuses.items())),
                    "observed_dates": sum(row["observed_at"] is not None for row in rows),
                    "sample_locator": json.loads(rows[0]["location"]) if rows else None,
                }
            )
        inventory = expected.get(path.name)
        files.append(
            {
                "path": f"data/BMW-Data/{path.name}",
                "kind": outcome["kind"],
                "bytes": len(data),
                "sha256": digest,
                "manifest_match": bool(inventory and inventory["sha256"] == digest and inventory["bytes"] == len(data)),
                "declared_scopes": declared,
                "surfaces": _surfaces(path, data, located),
                "scopes": scope_rows,
                "unplaced": outcome["unplaced"],
            }
        )
    missing_inputs = sorted(set(expected) - {Path(item["path"]).name for item in files})
    if set(first_files) != {Path(item["path"]).name for item in files}:
        raise AssertionError("Import result and local supported files differ")
    totals = {
        "files": len(files),
        "scopes": sum(len(item["scopes"]) for item in files),
        "parsed_records": sum(scope["parsed_records"] for item in files for scope in item["scopes"]),
        "stored_records": sum(scope["stored_records"] for item in files for scope in item["scopes"]),
        "same_source_duplicate_records": sum(
            scope["same_source_duplicate_records"] for item in files for scope in item["scopes"]
        ),
        "first_pass_new_evidence": first["totals"].get("new_evidence", 0),
        "repeat_new_evidence": repeat["totals"].get("new_evidence", 0),
    }
    if totals["repeat_new_evidence"] != 0 or not all(item["manifest_match"] for item in files):
        raise AssertionError("Repeat import or source manifest check failed")
    return {
        "method_version": METHOD,
        "import_method_version": "bmw-import-v1",
        "dedup_key": "original SHA-256 + vehicle + market",
        "totals": totals,
        "missing_manifest_inputs": missing_inputs,
        "annual_report_in_import": "BMW-Group-Report-2025-en.pdf" in {Path(item["path"]).name for item in files},
        "files": files,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--data-dir", required=True, help="Parent directory of BMW-Data")
    parser.add_argument("--db", required=True, help="Ignored SQLite database path")
    parser.add_argument("--manifest", default="notes/source-inventory.json")
    parser.add_argument("--output", help="JSON output path; otherwise stdout")
    args = parser.parse_args()
    root = Path(args.data_dir).resolve()
    folder = root / "BMW-Data"
    os.environ["DATA_DIR"] = str(root)
    os.environ["DB_PATH"] = str(Path(args.db).resolve())
    init_database()
    command = {"directory": "BMW-Data"}
    first = run_capability("market_data.import_bmw", command, "system:source-reconciliation")["outputs"]
    repeat = run_capability("market_data.import_bmw", command, "system:source-reconciliation")["outputs"]
    with closing(connect()) as connection:
        result = reconcile(folder, first, repeat, json.loads(Path(args.manifest).read_text()), connection)
    rendered = json.dumps(result, indent=2) + "\n"
    if args.output:
        Path(args.output).write_text(rendered)
    else:
        print(rendered, end="")


if __name__ == "__main__":
    main()
